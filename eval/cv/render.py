"""Renders the hand-written CV specs in eval/cv/specs into test documents.

Each spec is the ground truth; this script only lays it out in different
templates (single column, sidebar, tables, multi-page with repeated header and
footer, DOCX, rasterised scans, images). Output goes to eval/cv/fixtures with a
manifest of SHA-256 hashes so the TypeScript evaluation can detect stale files.

    services/resume-extractor/.venv/bin/python -I eval/cv/render.py
"""
import hashlib
import html
import os
import io
import json
import random
import sys
from pathlib import Path

import fitz
from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Pt
from PIL import Image, ImageFilter

ROOT = Path(__file__).resolve().parent
# CV_EVAL_SET=holdout renders the held-out set (eval/cv/holdout), which is never used for tuning.
SET_DIR = ROOT / 'holdout' if os.environ.get('CV_EVAL_SET') == 'holdout' else ROOT
SPECS = SET_DIR / 'specs'
OUT = SET_DIR / 'fixtures'
RENDER_VERSION = 1

MONTHS = ['January', 'February', 'March', 'April', 'May', 'June', 'July', 'August', 'September', 'October', 'November', 'December']
A4 = fitz.paper_rect('a4')


def fmt_date(value, style):
    if not value:
        return ''
    parts = value.split('-')
    year = parts[0]
    if len(parts) == 1 or style == 'year':
        return year
    m = int(parts[1])
    if style == 'month':
        return f'{MONTHS[m - 1]} {year}'
    if style == 'num':
        return f'{m:02d}/{year}'
    return f'{MONTHS[m - 1][:3]} {year}'


def date_range(item, spec, start_key='start', end_key='end'):
    if item.get('dateText'):
        return item['dateText']
    style, sep = spec.get('dateStyle', 'mon'), spec.get('rangeSep', ' – ')
    start = fmt_date(item.get(start_key), style)
    if item.get('endText'):
        end = item['endText']
    elif item.get('current') and not item.get(end_key):
        end = 'Present'
    elif item.get('current'):
        end = f"{fmt_date(item.get(end_key), style)} (Expected)"
    else:
        end = fmt_date(item.get(end_key), style)
    if start and end:
        return f'{start}{sep}{end}'
    return start or end


def lig(text, spec):
    return text.replace('fi', 'ﬁ').replace('fl', 'ﬂ') if spec.get('ligatures') else text


def esc(text, spec):
    return html.escape(lig(text, spec))


# ---------------------------------------------------------------- HTML (PDF)

CSS = '''
* { font-family: sans-serif; }
body { font-size: 9.5pt; }
h1 { font-size: 20pt; margin: 0 0 2pt 0; }
h2 { font-size: 11pt; margin: 9pt 0 3pt 0; text-transform: uppercase; border-bottom: 1px solid black; }
p { margin: 1pt 0; }
ul { margin: 1pt 0 3pt 14pt; padding: 0; }
li { margin: 0; }
table { width: 100%; border-collapse: collapse; }
td { padding: 0; vertical-align: top; }
td.r { text-align: right; }
.sub { font-style: italic; }
'''


def contact_html(spec):
    p = spec['person']
    icon = p.get('contactIcons')
    bits = []
    if p.get('email'):
        bits.append(('✉ ' if icon else '') + html.escape(p['email']))
    if p.get('phone'):
        bits.append(('☎ ' if icon else '') + html.escape(p['phone']))
    if p.get('location'):
        bits.append(html.escape(p['location']))
    links = [f'<a href="{html.escape(l["url"])}">{html.escape(l["display"])}</a>' for l in p.get('links', [])]
    return bits, links


RIGHT_TEXT = {}


def row(left, right, url=None):
    """A row whose right-hand plain text (dates, a link) is drawn right-aligned later.

    MuPDF's HTML engine ignores table widths and right-aligned cells, so the
    row is a paragraph with an id; its placed rectangle is recorded and the
    right-hand text is written at the column's right edge after layout.
    """
    if not right:
        return f'<p>{left}</p>'
    rid = f'rt{len(RIGHT_TEXT)}'
    RIGHT_TEXT[rid] = (right, url)
    return f'<p id="{rid}">{left}</p>'


def bullets(items, spec):
    if not items:
        return ''
    return '<ul>' + ''.join(f'<li>{esc(b, spec)}</li>' for b in items) + '</ul>'


def section_html(sec, spec, table_mode=False):
    kind, out = sec['kind'], [f'<h2>{html.escape(sec["heading"])}</h2>']
    if kind == 'education':
        if table_mode and sec.get('table'):
            out.append('<table border="1"><tr><td><b>Degree</b></td><td><b>Institution</b></td><td><b>Year</b></td><td><b>Score</b></td></tr>')
            for e in sec['items']:
                deg = e['degree'] + (f' {e["field"]}' if e.get('field') else '')
                out.append(f'<tr><td>{esc(deg, spec)}</td><td>{esc(e["institution"], spec)}</td><td>{html.escape(date_range(e, spec))}</td><td>{html.escape(e.get("grade", {}).get("text", ""))}</td></tr>')
            out.append('</table>')
        else:
            for e in sec['items']:
                inst = f'<b>{esc(e["institution"], spec)}</b>' + (f', {html.escape(e["location"])}' if e.get('location') else '')
                deg = esc(e['degree'], spec) + (f', {esc(e["field"], spec)}' if e.get('field') else '')
                if e.get('grade'):
                    deg += f'; {html.escape(e["grade"]["text"])}'
                out.append(row(inst, date_range(e, spec)))
                out.append(f'<p class="sub">{deg}</p>')
    elif kind in ('experience', 'volunteering'):
        for x in sec['items']:
            title = x.get('title') or x.get('role')
            org = x.get('organization')
            out.append(row(f'<b>{esc(title, spec)}</b>, {esc(org, spec)}', date_range(x, spec)))
            if x.get('location'):
                out.append(f'<p class="sub">{html.escape(x["location"])}</p>')
            out.append(bullets(x.get('bullets'), spec))
    elif kind == 'projects':
        for pr in sec['items']:
            head = f'<b>{esc(pr["name"], spec)}</b> | {html.escape(", ".join(pr["technologies"]))}'
            url = pr.get('url')
            out.append(row(head, url.replace('https://', '') if url else '', url))
            out.append(bullets(pr.get('bullets'), spec))
    elif kind == 'skills':
        if table_mode and sec.get('table'):
            out.append('<table border="1">' + ''.join(f'<tr><td><b>{html.escape(g.get("label", ""))}</b></td><td>{html.escape(", ".join(g["items"]))}</td></tr>' for g in sec['groups']) + '</table>')
        else:
            for g in sec['groups']:
                label = f'<b>{html.escape(g["label"])}:</b> ' if g.get('label') else ''
                out.append(f'<p>{label}{html.escape(", ".join(g["items"]))}</p>')
    elif kind == 'certifications':
        items = []
        for c in sec['items']:
            text = esc(c['name'], spec)
            if c.get('issuer'):
                text += f' — {html.escape(c["issuer"])}'
            when = c.get('dateText') or fmt_date(c.get('date'), spec.get('dateStyle', 'mon'))
            if when:
                text += f', {html.escape(when)}'
            if c.get('credentialId'):
                text += f' (Credential ID {html.escape(c["credentialId"])})'
            items.append(text)
        out.append('<ul>' + ''.join(f'<li>{t}</li>' for t in items) + '</ul>')
    elif kind == 'awards':
        out.append('<ul>' + ''.join(f'<li>{esc(a["title"], spec)}' + (f' ({html.escape(a["date"])})' if a.get('date') and a['date'] not in a['title'] else '') + '</li>' for a in sec['items']) + '</ul>')
    elif kind == 'publications':
        out.append('<ol>' + ''.join(f'<li>{esc(p["title"], spec)}. <i>{html.escape(p["venue"])}</i>, {html.escape(p["date"])}.</li>' for p in sec['items']) + '</ol>')
    elif kind == 'languages':
        out.append('<p>' + html.escape(', '.join(l['language'] + (f' ({l["proficiency"]})' if l.get('proficiency') else '') for l in sec['items'])) + '</p>')
    elif kind == 'memberships':
        out.append('<ul>' + ''.join(f'<li>{esc(m["title"], spec)}</li>' for m in sec['items']) + '</ul>')
    elif kind == 'references':
        for r in sec['items']:
            out.append(f'<p><b>{html.escape(r["name"])}</b>, {html.escape(r["relation"])}</p>')
            out.append(f'<p>{html.escape(" | ".join(v for v in (r.get("email"), r.get("phone")) if v))}</p>')
    return ''.join(out)


def header_html(spec, center=False):
    p = spec['person']
    bits, links = contact_html(spec)
    align = ' style="text-align:center"' if center else ''
    out = f'<h1{align}>{html.escape(p["name"])}</h1>'
    if p.get('headline'):
        out += f'<p{align}>{esc(p["headline"], spec)}</p>'
    out += f'<p{align}>{" | ".join(bits)}</p>'
    if links:
        out += f'<p{align}>{" | ".join(links)}</p>'
    return out


def summary_html(spec):
    if not spec.get('summary'):
        return ''
    return f'<h2>{html.escape(spec.get("summaryHeading", "Summary"))}</h2><p>{esc(spec["summary"], spec)}</p>'


def doc_html(spec, sections, table_mode=False, with_header=True):
    body = (header_html(spec) if with_header else '') + summary_html(spec)
    body += ''.join(section_html(s, spec, table_mode) for s in sections)
    return body


FONT = fitz.Font('helv')


def layout_pdf(path, columns, font_size, max_pages=1, header=None, footer=False, sidebar_fill=None):
    """Flows each (html, rect) column through pages, then draws the recorded
    right-hand texts and any running header/footer on every page."""
    stories = [(fitz.Story(body, user_css=CSS + f'body {{ font-size: {font_size}pt; }}'), rect) for body, rect in columns]
    placed = []
    buf = io.BytesIO()
    writer = fitz.DocumentWriter(buf)
    pending = list(range(len(stories)))
    page_no = 0
    while pending:
        page_no += 1
        if page_no > max_pages:
            raise SystemExit(f'{path.name}: content needs more than {max_pages} page(s)')
        dev = writer.begin_page(A4)
        still = []
        for i in pending:
            story, rect = stories[i]
            more, _ = story.place(rect)
            def record(p, pg=page_no, r=rect):
                if p.id in RIGHT_TEXT and p.open_close == 1:
                    placed.append((pg, p.id, p.rect, r))
            story.element_positions(lambda p: record(p))
            story.draw(dev)
            if more:
                still.append(i)
        writer.end_page()
        pending = still
    writer.close()
    doc = fitz.open('pdf', buf.getvalue())
    for page in doc:
        # The story's content stream leaves its flipped CTM active; isolate it.
        page.wrap_contents()
    for pg, rid, rect, col in placed:
        page = doc[pg - 1]
        text, url = RIGHT_TEXT[rid]
        tw = fitz.TextWriter(page.rect)
        width = FONT.text_length(text, fontsize=font_size)
        x0, baseline = col.x1 - width - 2, rect[1] + font_size * 0.98
        tw.append((x0, baseline), text, font=FONT, fontsize=font_size, )
        tw.write_text(page, color=(0.05, 0.3, 0.75) if url else None)
        if url:
            page.insert_link({'kind': fitz.LINK_URI, 'from': fitz.Rect(x0, baseline - font_size, x0 + width, baseline + 2), 'uri': url})
    for i, page in enumerate(doc, 1):
        if sidebar_fill:
            page.draw_rect(sidebar_fill, color=None, fill=(0.93, 0.95, 0.98), overlay=False)
        tw = fitz.TextWriter(page.rect)
        if header:
            tw.append((48, 40), header, font=FONT, fontsize=8.5)
        if footer:
            label = f'Page {i} of {doc.page_count}'
            tw.append(((A4.width - FONT.text_length(label, fontsize=8.5)) / 2, A4.height - 30), label, font=FONT, fontsize=8.5)
        tw.write_text(page)
    if max_pages > 1 and doc.page_count < 2:
        raise SystemExit(f'{path.name}: multipage variant fits on one page')
    doc.save(path, garbage=3, deflate=True)


def write_single(spec, path, table_mode=False):
    body = doc_html(spec, spec['sections'], table_mode)
    layout_pdf(path, [(body, fitz.Rect(40, 34, A4.width - 40, A4.height - 30))], 9.5)


def write_sidebar(spec, path):
    side = [s for s in spec['sections'] if s.get('sidebar')]
    main = [s for s in spec['sections'] if not s.get('sidebar')]
    p = spec['person']
    bits, links = contact_html(spec)
    side_html = '<h2>Contact</h2>' + ''.join(f'<p>{b}</p>' for b in bits) + ''.join(f'<p>{l}</p>' for l in links)
    side_html += ''.join(section_html(s, spec) for s in side)
    head = f'<h1>{html.escape(p["name"])}</h1>' + (f'<p>{esc(p["headline"], spec)}</p>' if p.get('headline') else '')
    main_html = head + summary_html(spec) + ''.join(section_html(s, spec) for s in main)
    layout_pdf(path, [(side_html, fitz.Rect(22, 34, 196, A4.height - 30)), (main_html, fitz.Rect(212, 34, A4.width - 30, A4.height - 30))], 9, sidebar_fill=fitz.Rect(0, 0, 204, A4.height))


def write_multipage(spec, path):
    body = doc_html(spec, spec['sections'])
    layout_pdf(path, [(body, fitz.Rect(48, 58, A4.width - 48, A4.height - 52))], 11, max_pages=4, header=spec.get('pageHeader', spec['person']['name']), footer=True)


# ---------------------------------------------------------------- DOCX

def add_hyperlink(paragraph, url, text):
    part = paragraph.part
    r_id = part.relate_to(url, 'http://schemas.openxmlformats.org/officeDocument/2006/relationships/hyperlink', is_external=True)
    link = OxmlElement('w:hyperlink')
    link.set(qn('r:id'), r_id)
    run = OxmlElement('w:r')
    t = OxmlElement('w:t')
    t.text = text
    run.append(t)
    link.append(run)
    paragraph._p.append(link)


def docx_section(container, sec, spec, table_mode):
    kind = sec['kind']
    style = spec.get('dateStyle', 'mon')
    h = container.add_paragraph(sec['heading'])
    h.style = 'Heading 2'
    def bullet(text):
        container.add_paragraph(text, style='List Bullet')
    if kind == 'education':
        if table_mode and sec.get('table'):
            t = container.add_table(rows=1, cols=4)
            for i, label in enumerate(('Degree', 'Institution', 'Year', 'Score')):
                t.rows[0].cells[i].text = label
            for e in sec['items']:
                cells = t.add_row().cells
                cells[0].text = e['degree'] + (f' {e["field"]}' if e.get('field') else '')
                cells[1].text = e['institution']
                cells[2].text = date_range(e, spec)
                cells[3].text = e.get('grade', {}).get('text', '')
        else:
            for e in sec['items']:
                p = container.add_paragraph()
                p.add_run(e['institution']).bold = True
                if e.get('location'):
                    p.add_run(f', {e["location"]}')
                p.add_run(f'\t{date_range(e, spec)}')
                line = e['degree'] + (f', {e["field"]}' if e.get('field') else '') + (f'; {e["grade"]["text"]}' if e.get('grade') else '')
                container.add_paragraph(line).runs[0].italic = True
    elif kind in ('experience', 'volunteering'):
        for x in sec['items']:
            p = container.add_paragraph()
            p.add_run(x.get('title') or x.get('role')).bold = True
            p.add_run(f' at {x["organization"]}')
            dr = date_range(x, spec)
            if dr:
                p.add_run(f'\t{dr}')
            if x.get('location'):
                container.add_paragraph(x['location']).runs[0].italic = True
            for b in x.get('bullets', []):
                bullet(b)
    elif kind == 'projects':
        for pr in sec['items']:
            p = container.add_paragraph()
            p.add_run(pr['name']).bold = True
            p.add_run(f' ({", ".join(pr["technologies"])})')
            if pr.get('url'):
                p.add_run(' ')
                add_hyperlink(p, pr['url'], pr['url'].replace('https://', ''))
            for b in pr.get('bullets', []):
                bullet(b)
    elif kind == 'skills':
        if table_mode and sec.get('table'):
            t = container.add_table(rows=0, cols=2)
            for g in sec['groups']:
                cells = t.add_row().cells
                cells[0].text = g.get('label', '')
                cells[1].text = ', '.join(g['items'])
        else:
            for g in sec['groups']:
                container.add_paragraph((f'{g["label"]}: ' if g.get('label') else '') + ', '.join(g['items']))
    elif kind == 'certifications':
        for c in sec['items']:
            text = c['name'] + (f' — {c["issuer"]}' if c.get('issuer') else '')
            when = c.get('dateText') or fmt_date(c.get('date'), style)
            bullet(text + (f', {when}' if when else '') + (f' (Credential ID {c["credentialId"]})' if c.get('credentialId') else ''))
    elif kind == 'awards':
        for a in sec['items']:
            bullet(a['title'] + (f' ({a["date"]})' if a.get('date') and a['date'] not in a['title'] else ''))
    elif kind == 'publications':
        for i, p in enumerate(sec['items'], 1):
            container.add_paragraph(f'{i}. {p["title"]}. {p["venue"]}, {p["date"]}.')
    elif kind == 'languages':
        container.add_paragraph(', '.join(l['language'] + (f' ({l["proficiency"]})' if l.get('proficiency') else '') for l in sec['items']))
    elif kind == 'memberships':
        for m in sec['items']:
            bullet(m['title'])
    elif kind == 'references':
        for r in sec['items']:
            container.add_paragraph(f'{r["name"]}, {r["relation"]}')
            container.add_paragraph(' | '.join(v for v in (r.get('email'), r.get('phone')) if v))


def write_docx(spec, path, table_mode=False):
    d = Document()
    d.styles['Normal'].font.size = Pt(10)
    p = spec['person']
    contact = [v for v in (p.get('email'), p.get('phone'), p.get('location')) if v]
    header = d.sections[0].header
    if spec.get('pageHeader'):
        header.paragraphs[0].text = spec['pageHeader']
    if spec.get('docxContactIn') == 'header':
        hp = header.paragraphs[0]
        hp.text = ' | '.join(contact)
        hp.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    title = d.add_paragraph()
    title.add_run(p['name']).bold = True
    title.runs[0].font.size = Pt(20)
    if p.get('headline'):
        d.add_paragraph(p['headline'])
    if spec.get('docxContactIn') != 'header':
        d.add_paragraph(' | '.join(contact))
    if p.get('links'):
        lp = d.add_paragraph()
        for i, l in enumerate(p['links']):
            if i:
                lp.add_run(' | ')
            add_hyperlink(lp, l['url'], l['display'])
    if spec.get('summary'):
        d.add_paragraph(spec.get('summaryHeading', 'Summary')).style = 'Heading 2'
        d.add_paragraph(spec['summary'])
    if table_mode:
        # Two-column layout table: sidebar sections left, main sections right.
        side = [s for s in spec['sections'] if s['kind'] in ('skills', 'languages', 'certifications', 'memberships')]
        main = [s for s in spec['sections'] if s not in side]
        layout = d.add_table(rows=1, cols=2)
        left, right = layout.rows[0].cells
        for s in side:
            docx_section(left, s, spec, True)
        for s in main:
            docx_section(right, s, spec, True)
    else:
        for s in spec['sections']:
            docx_section(d, s, spec, False)
    if spec.get('pageHeader'):
        d.sections[0].footer.paragraphs[0].text = 'Curriculum Vitae'
    d.save(path)


# ---------------------------------------------------------------- scans and images

def rasterise(pdf_path, dpi):
    with fitz.open(pdf_path) as doc:
        return [Image.open(io.BytesIO(p.get_pixmap(dpi=dpi, alpha=False).tobytes('png'))).convert('RGB') for p in doc]


def degrade(im, seed):
    rnd = random.Random(seed)
    im = im.rotate(1.4, resample=Image.Resampling.BICUBIC, expand=True, fillcolor=(255, 255, 255))
    im = im.filter(ImageFilter.GaussianBlur(radius=0.7))
    px = im.load()
    for _ in range(im.width * im.height // 180):
        x, y = rnd.randrange(im.width), rnd.randrange(im.height)
        v = rnd.randrange(0, 140)
        px[x, y] = (v, v, v)
    buf = io.BytesIO()
    im.save(buf, 'JPEG', quality=38)
    return Image.open(io.BytesIO(buf.getvalue())).convert('RGB')


def images_to_pdf(images, path, dpi):
    doc = fitz.open()
    for im in images:
        w, h = im.width * 72 / dpi, im.height * 72 / dpi
        page = doc.new_page(width=w, height=h)
        buf = io.BytesIO()
        im.save(buf, 'JPEG', quality=85)
        page.insert_image(page.rect, stream=buf.getvalue())
    doc.save(path, garbage=3, deflate=True)


def main():
    OUT.mkdir(exist_ok=True)
    manifest = {'renderVersion': RENDER_VERSION, 'documents': []}
    only = set(sys.argv[1:])
    for spec_path in sorted(SPECS.glob('*.json')):
        spec = json.loads(spec_path.read_text())
        if only and spec['id'] not in only:
            continue
        primary_pdf = None
        for variant in spec['variants']:
            ext = {'docx-plain': 'docx', 'docx-table': 'docx', 'png': 'png', 'jpg': 'jpg'}.get(variant, 'pdf')
            path = OUT / f'{spec["id"]}--{variant}.{ext}'
            if variant == 'pdf-single':
                write_single(spec, path)
            elif variant == 'pdf-table':
                write_single(spec, path, table_mode=True)
            elif variant == 'pdf-sidebar':
                write_sidebar(spec, path)
            elif variant == 'pdf-multipage':
                write_multipage(spec, path)
            elif variant == 'docx-plain':
                write_docx(spec, path)
            elif variant == 'docx-table':
                write_docx(spec, path, table_mode=True)
            elif variant in ('scan-good', 'scan-poor', 'png', 'jpg'):
                source = primary_pdf or OUT / f'{spec["id"]}--_source.pdf'
                if primary_pdf is None:
                    write_single(spec, source)
                if variant == 'scan-good':
                    images_to_pdf(rasterise(source, 200), path, 200)
                elif variant == 'scan-poor':
                    images_to_pdf([degrade(im, spec['id']) for im in rasterise(source, 110)], path, 110)
                else:
                    im = rasterise(source, 150)[0]
                    im.save(path, 'PNG' if variant == 'png' else 'JPEG', **({} if variant == 'png' else {'quality': 82}))
                if source.name.endswith('--_source.pdf'):
                    source.unlink()
            else:
                raise SystemExit(f'unknown variant {variant}')
            if variant.startswith('pdf-') and primary_pdf is None:
                primary_pdf = path
            manifest['documents'].append({
                'file': path.name,
                'spec': spec['id'],
                'variant': variant,
                'sha256': hashlib.sha256(path.read_bytes()).hexdigest(),
            })
            print(f'wrote {path.name}')
    if not only:
        (OUT / 'manifest.json').write_text(json.dumps(manifest, indent=2) + '\n')


if __name__ == '__main__':
    main()
