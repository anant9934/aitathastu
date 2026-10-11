"""Create synthetic CV PDFs for real raster and vector-outline OCR checks.

All content is fictional. Outputs stay in the supplied temporary directory.
"""
import os
import sys
from PIL import Image, ImageDraw, ImageFont
from reportlab.pdfgen import canvas
from reportlab.lib.utils import ImageReader
from docx import Document

destination = sys.argv[1]
os.makedirs(destination, exist_ok=True)
width, height, scale = 612, 792, 2
image = Image.new("L", (width * scale, height * scale), 255)
draw = ImageDraw.Draw(image)
font_path = "/System/Library/Fonts/Supplemental/Arial.ttf"
heading = ImageFont.truetype(font_path, 42)
body = ImageFont.truetype(font_path, 29)
lines = [
    ("ALEX MORGAN", heading),
    ("Email: alex@example.test", body),
    ("EDUCATION", heading),
    ("Example University, BSc Computer Science", body),
    ("EXPERIENCE", heading),
    ("Research Intern at Example Lab, 2024", body),
    ("SKILLS", heading),
    ("Python, PyTorch, SQL", body),
    ("PROJECTS", heading),
    ("Garden Vision image classifier", body),
    ("https://github.com/alexexample", body),
    ("https://linkedin.com/in/alexexample", body),
]
y = 90
for text, font in lines:
    draw.text((90, y), text, font=font, fill=0)
    y += 116 if font is heading else 100

image_path = os.path.join(destination, "synthetic_image_cv.pdf")
c = canvas.Canvas(image_path, pagesize=(width, height))
c.drawImage(ImageReader(image.convert("RGB")), 0, 0, width, height)
c.save()

native_path = os.path.join(destination, "synthetic_native_cv.pdf")
c = canvas.Canvas(native_path, pagesize=(width, height))
y_native = 740
for line, font in lines:
    c.setFont("Helvetica-Bold" if font is heading else "Helvetica", 20 if font is heading else 14)
    c.drawString(45, y_native, line)
    y_native -= 55
c.save()

vector_path = os.path.join(destination, "synthetic_vector_cv.pdf")
c = canvas.Canvas(vector_path, pagesize=(width, height), pageCompression=1)
c.setFillColorRGB(0, 0, 0)
pixels = image.load()
for row in range(image.height):
    start = None
    for col in range(image.width + 1):
        dark = col < image.width and pixels[col, row] < 128
        if dark and start is None:
            start = col
        elif not dark and start is not None:
            c.rect(start / scale, height - (row + 1) / scale,
                   (col - start) / scale, 1 / scale, fill=1, stroke=0)
            start = None
c.save()

print(image_path)
print(vector_path)
print(native_path)

document = Document()
document.add_heading("ALEX MORGAN", 0)
document.add_paragraph("Email: alex@example.test")
document.add_heading("EDUCATION", 1)
table = document.add_table(rows=1, cols=2)
table.cell(0, 0).text = "Example University"
table.cell(0, 1).text = "BSc Computer Science"
document.add_heading("EXPERIENCE", 1)
document.add_paragraph("Research Intern at Example Lab, 2024")
document.add_heading("SKILLS", 1)
document.add_paragraph("Python, PyTorch, SQL")
document.add_paragraph("https://github.com/alexexample")
document.add_paragraph("https://linkedin.com/in/alexexample")
docx_path = os.path.join(destination, "synthetic_table_cv.docx")
document.save(docx_path)
print(docx_path)

standard_docx = Document()
standard_docx.add_heading("ALEX MORGAN", 0)
standard_docx.add_paragraph("Example University, BSc Computer Science")
standard_docx.add_paragraph("Research Intern at Example Lab")
standard_docx.add_paragraph("Python, SQL, PyTorch")
standard_docx.save(os.path.join(destination, "synthetic_standard_cv.docx"))

mixed_path = os.path.join(destination, "synthetic_mixed_cv.pdf")
c = canvas.Canvas(mixed_path, pagesize=(width, height))
c.setFont("Helvetica-Bold", 20)
c.drawString(45, 740, "ALEX MORGAN")
c.drawString(45, 690, "EDUCATION")
c.setFont("Helvetica", 14)
c.drawString(45, 650, "Example University, BSc Computer Science")
c.showPage()
c.drawImage(ImageReader(image.convert("RGB")), 0, 0, width, height)
c.save()

columns_path = os.path.join(destination, "synthetic_two_column_cv.pdf")
c = canvas.Canvas(columns_path, pagesize=(width, height))
c.setFont("Helvetica-Bold", 20)
c.drawString(45, 745, "ALEX MORGAN")
c.setFont("Helvetica-Bold", 15)
c.drawString(45, 685, "EDUCATION")
c.drawString(325, 685, "SKILLS")
c.setFont("Helvetica", 12)
c.drawString(45, 650, "Example University")
c.drawString(45, 630, "BSc Computer Science")
c.drawString(325, 650, "Python, SQL, PyTorch")
c.save()

short_path = os.path.join(destination, "synthetic_short_cv.pdf")
c = canvas.Canvas(short_path, pagesize=(width, height))
c.setFont("Helvetica", 14)
c.drawString(45, 740, "ALEX MORGAN")
c.drawString(45, 710, "SKILLS")
c.drawString(45, 680, "Python")
c.save()

with open(os.path.join(destination, "synthetic_corrupted.pdf"), "wb") as output:
    output.write(b"%PDF-1.7\ncorrupted and truncated")
